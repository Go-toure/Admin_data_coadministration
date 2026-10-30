#!/usr/bin/env python3
"""
admin.py - one command for the whole AFRO admin data cycle.

    python admin.py status            where everything stands (runs, backups, git, folders)
    python admin.py backup            save the live SharePoint file locally (and list the versions)
    python admin.py build             build the next version from the LIVE file (backup first)
    python admin.py check             verify the newest build against the file it started from
    python admin.py excel             store the check_dp values in the build (written directly; --use-excel to use Excel)
    python admin.py publish           replace the file on SharePoint (all gates must be green)
    python admin.py all               backup -> build -> check -> excel -> publish
    python admin.py runs              list the builds
    python admin.py inspect           quality + hygiene + lookup checks on a file (the live file by default)
    python admin.py find ...          look up records (country, province, district, response, date, round) in any version
    python admin.py compare A B       what differs between two versions (builds v24, backups b9, files)
    python admin.py rules             what the system knows and fixes, and the decisions behind it
    python admin.py commit -m "msg" FILE [FILE ...]    git add + commit + push (code only, never data)
    python admin.py sync-scope        copy the scripts / the published file to the scope folders
    python admin.py doctor            check python packages, folders, Excel

Every build lives in its own folder, workspace/runs/vNN, with a manifest.json that records what it started from,
what was done and what is still to do. publish refuses to run unless the build was checked, saved in Excel and the
live file is still the one the build started from (so nobody's newer edit is ever overwritten).

Settings: admin_settings.json (paths only, no secrets). The SharePoint credentials stay in the backup tool's
config/secrets.env, which this script never reads.
"""
import argparse
import csv
import datetime as dt
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SETTINGS_FILE = HERE / "admin_settings.json"
DEFAULTS = {
    "backup_repo": "../SharePoint_version_backup",
    "workspace": "workspace",
    "version_floor": 23,                      # the next build is max(floor, existing runs) + 1
    "builder": "add_coadmin_columns.py",
    "publisher": "publish_admin_data.py",
    "lookup": "lookup.xlsx",
    "geo": "geo_reference.csv",
    "district_source": "",                    # optional recovered file used by --district-source
    "build_extra_args": [],
    "max_row_drop_pct": 2.0,
    "key_columns": ["Country", "Province", "District", "SIA_date", "Round_Add", "Vaccine_type", "Response"],
    "scope_script_dir": "",
    "scope_data_dir": "",
    "excel_timeout_min": 30,
    "excel_method": "python",                 # "python": values written into the file, no Excel needed; "excel": desktop Excel
    "excel_visible": True,                    # show the Excel window while it saves (any pop-up can then be answered)
}
EXPECTED_CHANGE = {"Response", "Round_Add", "District", "Province", "Country", "SIA_date", "Vaccine_type", "PopPolio",
                   "CVPolio", "TotalNbrVaccPolio", "Tot_Vacc_nOPV2", "CV_nOPV2", "Total_Vacc_bOPV", "CV_bOPV", "TotPop",
                   "AgeGPolio", "Admin_3", "Year", "Index", "SIAMTH"}
DATA_EXT = (".xlsx", ".xls", ".csv", ".pkl", ".rds", ".zip")
DATA_OK = ("lookup.xlsx", "geo_reference.csv")
FORBIDDEN_MSG = ("co-authored-by", "claude-session")


# ----------------------------------------------------------------------------------------------- settings / logging
def load_settings():
    s = dict(DEFAULTS)
    if SETTINGS_FILE.is_file():
        try:
            s.update(json.loads(SETTINGS_FILE.read_text(encoding="utf-8")))
        except Exception as e:
            die("admin_settings.json is not valid JSON: %s" % e)
    return s


S = load_settings()


def P(rel):
    """A path from the settings: absolute stays, relative is taken from this folder."""
    p = Path(str(rel))
    return p if p.is_absolute() else (HERE / p).resolve()


WORK = P(S["workspace"])
RUNS = WORK / "runs"
STATE = WORK / "state.json"
LOGDIR = HERE / "logs"
BACKUP_REPO = P(S["backup_repo"])


def log(msg="", quiet=False):
    line = "%s  %s" % (dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), msg) if msg else ""
    if not quiet:
        print(line, flush=True)
    try:
        LOGDIR.mkdir(exist_ok=True)
        with open(LOGDIR / ("admin_%s.log" % dt.date.today().strftime("%Y%m%d")), "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass


def die(msg, code=2):
    print("\nSTOPPED: " + msg + "\n", flush=True)
    try:
        log("STOPPED: " + msg, quiet=True)
    except Exception:
        pass
    sys.exit(code)


def run_cmd(cmd, cwd=None, tee=None, check=False):
    """Run a command, stream its output, optionally copy it to a file. Returns (returncode, text)."""
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUNBUFFERED="1")
    p = subprocess.Popen([str(c) for c in cmd], cwd=str(cwd) if cwd else None, stdout=subprocess.PIPE,
                         stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", env=env)
    out = []
    f = open(tee, "w", encoding="utf-8") if tee else None
    try:
        for line in p.stdout:
            out.append(line)
            print(line, end="", flush=True)
            if f:
                f.write(line)
    finally:
        if f:
            f.close()
    p.wait()
    if check and p.returncode:
        die("a step failed (exit code %s): %s" % (p.returncode, " ".join(str(c) for c in cmd[:3])))
    return p.returncode, "".join(out)


# ----------------------------------------------------------------------------------------------- state and manifests
def read_json(p, default=None):
    try:
        return json.loads(Path(p).read_text(encoding="utf-8"))
    except Exception:
        return default


def write_json(p, obj):
    Path(p).parent.mkdir(parents=True, exist_ok=True)
    Path(p).write_text(json.dumps(obj, indent=2, ensure_ascii=False, default=str), encoding="utf-8")


def run_dirs():
    out = []
    if RUNS.is_dir():
        for d in RUNS.iterdir():
            m = re.match(r"^v(\d+)$", d.name)
            if m and d.is_dir():
                out.append((int(m.group(1)), d))
    return sorted(out)


def next_version():
    nums = [n for n, _ in run_dirs()] + [int(S["version_floor"])]
    return max(nums) + 1


def latest_run(unpublished=False):
    for n, d in reversed(run_dirs()):
        m = read_json(d / "manifest.json", {})
        if unpublished and m.get("published"):
            continue
        return n, d, m
    return None, None, None


def save_manifest(d, m):
    write_json(Path(d) / "manifest.json", m)


def now():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ----------------------------------------------------------------------------------------------- the SharePoint backup tool
def catalog_rows(key="admin"):
    p = BACKUP_REPO / "catalog" / "versions.csv"
    if not p.is_file():
        return []
    with open(p, newline="", encoding="utf-8") as f:
        rows = [r for r in csv.DictReader(f) if r.get("key") == key]
    return sorted(rows, key=lambda r: int(r["version"]))


def latest_snapshot():
    rows = catalog_rows("admin")
    if not rows:
        return None, None
    r = rows[-1]
    return r, BACKUP_REPO / r["snapshot"]


def do_backup(show=True):
    tool = BACKUP_REPO / "sp_versions.py"
    if not tool.is_file():
        die("the backup tool was not found: %s (see backup_repo in admin_settings.json)" % tool)
    log("Backing up the live SharePoint file ...")
    rc, _ = run_cmd([sys.executable, tool, "backup", "admin"], cwd=BACKUP_REPO)
    if rc:
        die("the SharePoint backup failed - nothing was changed. Check the message above.")
    if show:
        run_cmd([sys.executable, tool, "list", "admin"], cwd=BACKUP_REPO)
    return latest_snapshot()


# ----------------------------------------------------------------------------------------------- data helpers
def _pd():
    import pandas as pd
    return pd


def load_norm(path, sheet=0):
    """The first sheet as text, numbers written the same way whatever Excel did to them, text without the control
    characters / extra spaces that Excel or a browser add, without check_dp."""
    pd = _pd()
    df = pd.read_excel(path, sheet_name=sheet, dtype=str, keep_default_na=False, na_values=[""])
    if "check_dp" in df.columns:
        df = df.drop(columns=["check_dp"])
    for c in df.columns:
        col = df[c]
        num = pd.to_numeric(col, errors="coerce")
        ok = num.notna()
        if ok.any():
            df[c] = col.where(~ok, num.map(lambda x: format(x, ".10g") if x == x else x))
        col = df[c]
        dirty = col.str.contains(r"_x000D_|\s\s|^\s|\s$|[\r\n\t]", regex=True, na=False)
        if dirty.any():
            df.loc[dirty, c] = col[dirty].str.replace(r"_x000D_|\s+", " ", regex=True).str.strip()
    return df.fillna("")


def fingerprint(df):
    pd = _pd()
    h = hashlib.sha256()
    h.update("|".join(map(str, df.columns)).encode())
    h.update(pd.util.hash_pandas_object(df, index=False).values.tobytes())
    return h.hexdigest()


def import_builder():
    """The builder's own helper functions (lookup reading, country keys) are reused, never copied."""
    path = P(S["builder"])
    spec = importlib.util.spec_from_file_location("coadmin_builder", path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


# ----------------------------------------------------------------------------------------------- status
def git(*args, cwd=HERE):
    try:
        r = subprocess.run(["git", *args], cwd=str(cwd), capture_output=True, text=True, encoding="utf-8", errors="replace")
        return r.returncode, (r.stdout + r.stderr).strip()
    except FileNotFoundError:
        return 127, "git is not installed"


def cmd_status(a):
    print("\n=== AFRO admin data: status ===\n")
    print("Folder    : %s" % HERE)
    print("Workspace : %s" % WORK)
    print("Backups   : %s%s" % (BACKUP_REPO, "" if BACKUP_REPO.is_dir() else "   <-- NOT FOUND"))
    if a.live:
        do_backup(show=False)
    rows = catalog_rows("admin")
    if rows:
        r = rows[-1]
        print("\nLatest local backup of the live file: version %s, %s rows, saved %s UTC, by %s"
              % (r["version"], r["rows"], r["saved_at_utc"].replace("T", " ").replace("Z", ""), r["modified_by"]))
        if not a.live:
            print("  (run 'status --live' or 'backup' to read SharePoint now)")
    else:
        print("\nNo local backup yet: run 'backup'.")
    print("\nBuilds:")
    runs = run_dirs()
    if not runs:
        print("  none yet: run 'build'")
    for n, d in runs[-5:]:
        m = read_json(d / "manifest.json", {})
        print("  v%-3d %s" % (n, describe_run(m)))
    st = read_json(STATE, {})
    if st.get("published"):
        print("\nLast published: v%s on %s (SharePoint backup version %s)"
              % (st["published"].get("version"), st["published"].get("at"), st["published"].get("backup_version")))
    n, d, m = latest_run(unpublished=True)
    if n:
        print("\nNext step for v%d: %s" % (n, next_step(m)))
    else:
        print("\nNext step: build")
    print("\nGit:")
    rc, out = git("status", "--short")
    if rc:
        print("  (not available here: %s)" % out.splitlines()[0][:80])
    else:
        changed = [l for l in out.splitlines() if l.strip()]
        print("  uncommitted: %s" % (("%d item(s)" % len(changed)) if changed else "nothing"))
        for l in changed[:8]:
            print("    " + l)
        rc, out = git("log", "@{u}..HEAD", "--oneline")
        if rc == 0:
            print("  not pushed : %s" % ("%d commit(s)" % len(out.splitlines()) if out.strip() else "nothing"))
    left = sorted(p.name for p in HERE.glob("_history*") if p.is_dir())
    if left:
        print("\nLeftover folders (delete when their commit is done): %s" % ", ".join(left))
        print("  rm -r " + " ".join(left))
    print()


def describe_run(m):
    bits = []
    if m.get("build", {}).get("ok"):
        bits.append("built")
    c = m.get("check")
    if c:
        bits.append("check %s" % ("OK" if c.get("ok") else "FAILED"))
    if m.get("excel", {}).get("ok"):
        bits.append("saved in Excel")
    if m.get("published"):
        bits.append("PUBLISHED %s" % m["published"].get("at", ""))
    return ", ".join(bits) or "started"


def next_step(m):
    if not m.get("build", {}).get("ok"):
        return "build (the last one did not finish)"
    if not m.get("check"):
        return "check"
    if not m["check"].get("ok"):
        return "fix the problems in the check report, then build again"
    if not m.get("excel", {}).get("ok"):
        return "excel (open and save once)"
    return "publish"


def cmd_runs(a):
    for n, d in run_dirs():
        m = read_json(d / "manifest.json", {})
        b = m.get("base", {})
        print("v%-3d %-19s base: backup v%s (%s rows) | rows out %s | errors %s | %s"
              % (n, m.get("created", ""), b.get("backup_version", "?"), b.get("rows", "?"),
                 m.get("build", {}).get("rows_out", "?"), m.get("build", {}).get("errors", "?"), describe_run(m)))


# ----------------------------------------------------------------------------------------------- backup
def cmd_backup(a):
    do_backup(show=True)


# ----------------------------------------------------------------------------------------------- build
def cmd_build(a):
    for need in (P(S["builder"]), P(S["lookup"])):
        if not need.is_file():
            die("file not found: %s" % need)
    if a.base:
        base = Path(a.base)
        if not base.is_file():
            die("the base file does not exist: %s" % base)
        row = None
    else:
        row, base = (do_backup(show=False) if not a.no_backup else latest_snapshot())
        if row is None or not base.is_file():
            die("no local backup of the live file: run 'backup' first.")
    n = next_version()
    rd = RUNS / ("v%d" % n)
    if rd.exists():
        die("%s already exists." % rd)
    rd.mkdir(parents=True)
    out = rd / "out"
    log("Build v%d, starting from %s" % (n, base.name))
    log("Reading the base file (for the check later) ...")
    base_df = load_norm(base)
    m = {"version": n, "created": now(),
         "base": {"path": str(base), "backup_version": row["version"] if row else None,
                  "sha256": row["sha256"] if row else None, "rows": len(base_df), "columns": list(base_df.columns),
                  "fingerprint": fingerprint(base_df), "modified_by": row["modified_by"] if row else None},
         "build": {}}
    del base_df
    save_manifest(rd, m)
    cmd = [sys.executable, P(S["builder"]), base, "--out-dir", out, "--lookup", P(S["lookup"])]
    geo = P(S["geo"])
    if geo.is_file():
        cmd += ["--geo", geo]
    snaps = BACKUP_REPO / "snapshots" / "admin"
    if snaps.is_dir():
        cmd += ["--backups", snaps]
    ds = S.get("district_source")
    if ds and P(ds).is_file():
        cmd += ["--district-source", P(ds)]
    cmd += list(S.get("build_extra_args") or []) + list(a.extra or [])
    m["build"]["command"] = " ".join(str(c) for c in cmd)
    rc, text = run_cmd(cmd, cwd=HERE, tee=rd / "build.log")
    if rc:
        m["build"]["ok"] = False
        save_manifest(rd, m)
        die("the build failed (see %s)." % (rd / "build.log"))
    f_up = out / "upload_to_sharepoint" / "AFRO_admin_data.xlsx"
    if not f_up.is_file():
        die("the build produced no AFRO_admin_data.xlsx (see %s)." % (rd / "build.log"))
    # versioned names in the run folder (what the publish step and the team look at)
    final = rd / ("AFRO_admin_data_v%d.xlsx" % n)
    shutil.copy2(f_up, final)
    for src, name in ((out / "upload_to_sharepoint" / "AFRO_admin_data_dictionary.xlsx", "AFRO_admin_data_dictionary_v%d.xlsx" % n),
                      (out / "review" / "AFRO_admin_data_coadmin_audit.xlsx", "AFRO_admin_data_coadmin_audit_v%d.xlsx" % n),
                      (out / "review" / "AFRO_admin_data_quality_report.xlsx", "AFRO_admin_data_quality_report_v%d.xlsx" % n),
                      (out / "review" / "Feedback_suspicious_coverage.xlsx", "Feedback_suspicious_coverage_v%d.xlsx" % n)):
        if src.is_file():
            shutil.copy2(src, rd / name)
    mm = re.search(r"\((\d+) records with an Error", text) or re.search(r"(\d+) records with an Error", text)
    mr = re.search(r"rows out: (\d+)", text)
    m["build"].update({"ok": True, "finished": now(), "file": str(final), "errors": int(mm.group(1)) if mm else None,
                       "rows_out": int(mr.group(1)) if mr else None,
                       "summary": [l.strip() for l in text.splitlines()
                                   if re.search(r"filled|renamed|corrected|removed|merged|restored|repair", l)][:25]})
    save_manifest(rd, m)
    log("Build v%d done: %s" % (n, final))
    log("Next: python admin.py check")


# ----------------------------------------------------------------------------------------------- check
def response_country_mismatches(df, builder, lk):
    """Rows whose Response is a campaign name of the lookup that belongs to another country."""
    pd = _pd()
    owner = {}
    for rk_, ck_ in zip(lk["_rk"], lk["_ck"]):
        owner.setdefault(rk_, set()).add(ck_)
    pair = df[["Country", "Response"]].drop_duplicates()
    bad = {}
    for c, r in zip(pair["Country"], pair["Response"]):
        rk_ = builder._rk(r)
        if rk_ in owner and builder._lkc(c) not in owner[rk_]:
            bad[(c, r)] = int(((df["Country"] == c) & (df["Response"] == r)).sum())
    return bad


def cmd_check(a):
    n, rd, m = (None, None, None)
    if a.run:
        rd = RUNS / ("v%d" % a.run)
        n, m = a.run, read_json(rd / "manifest.json", {})
    else:
        n, rd, m = latest_run()
    if not n or not m.get("build", {}).get("ok"):
        die("no finished build to check: run 'build' first.")
    final = Path(m["build"]["file"])
    log("Checking v%d (%s) against the file it started from ..." % (n, final.name))
    pd = _pd()
    base_path = Path(m["base"]["path"])
    base = load_norm(base_path)
    new = load_norm(final)
    builder = import_builder()
    lk = builder.load_lookup(str(P(S["lookup"])))
    hard, warn, info = [], [], []

    # 1. header
    miss = [c for c in base.columns if c not in new.columns]
    added = [c for c in new.columns if c not in base.columns]
    if miss:
        hard.append("columns missing from the new file: %s" % ", ".join(miss[:10]))
    if added:
        warn.append("new columns: %s" % ", ".join(added[:10]))
    # 2. rows
    drop_pct = (len(base) - len(new)) / max(len(base), 1) * 100
    info.append("rows: %d -> %d (%+d)" % (len(base), len(new), len(new) - len(base)))
    if drop_pct > float(S["max_row_drop_pct"]):
        hard.append("%.1f%% of the rows disappeared (limit %.1f%%)" % (drop_pct, float(S["max_row_drop_pct"])))
    # 3. key fields blank
    keys = [k for k in S["key_columns"] if k in new.columns and k in base.columns]

    def blanks(df):
        return int((df[keys].apply(lambda c: c.astype(str).str.strip() == "")).any(axis=1).sum())
    b0, b1 = blanks(base), blanks(new)
    info.append("records with an empty key field: %d -> %d" % (b0, b1))
    if b1 > b0:
        hard.append("more records with an empty key field than before (%d -> %d)" % (b0, b1))
    # 4. Response that belongs to another country
    if "Response" in new.columns and "Country" in new.columns:
        m0 = response_country_mismatches(base, builder, lk)
        m1 = response_country_mismatches(new, builder, lk)
        info.append("records with a Response of another country (per the lookup): %d -> %d"
                    % (sum(m0.values()), sum(m1.values())))
        if sum(m1.values()) > sum(m0.values()):
            hard.append("more records with another country's Response (%d -> %d)" % (sum(m0.values()), sum(m1.values())))
        elif m1:
            warn.append("still %d record(s) with another country's Response: %s"
                        % (sum(m1.values()), "; ".join("%s / %s (%d)" % (k[0], k[1], v) for k, v in list(m1.items())[:6])))
    # 5. duplicates on the portal key
    if keys:
        d0 = int(base.duplicated(keys, keep=False).sum())
        d1 = int(new.duplicated(keys, keep=False).sum())
        info.append("records sharing the portal key with another record: %d -> %d" % (d0, d1))
        if d1 > d0:
            warn.append("more duplicate keys than before (%d -> %d)" % (d0, d1))
    # 6. errors reported by the build
    prev = None
    for k, d in reversed(run_dirs()):
        if k < n:
            e = read_json(d / "manifest.json", {}).get("build", {}).get("errors")
            if e is not None:
                prev = e
                break
    e1 = m["build"].get("errors")
    info.append("quality errors: %s -> %s" % (prev if prev is not None else "?", e1 if e1 is not None else "?"))
    if prev is not None and e1 is not None and e1 > prev:
        warn.append("more quality errors than the previous build (%d -> %d)" % (prev, e1))
    # 7. what changed, cell by cell (only when the rows line up)
    changes = {}
    if len(base) == len(new):
        common = [c for c in new.columns if c in base.columns]
        for c in common:
            k = int((base[c].values != new[c].values).sum())
            if k:
                changes[c] = k
        info.append("cells changed vs the base: %s" % (", ".join("%s %d" % kv for kv in sorted(changes.items(), key=lambda x: -x[1])[:12]) or "none"))
        big = [c for c in ("TotalNbrVaccPolio", "PopPolio") if changes.get(c, 0) > 0.01 * len(new)]
        if big:
            warn.append("more than 1%% of the records changed in: %s" % ", ".join(big))
        odd = [c for c, k in changes.items() if c not in EXPECTED_CHANGE and k > 0.005 * len(new)]
        if odd:
            warn.append("columns the builder does not normally touch changed in more than 0.5%% of the records: %s "
                        "(look at 'compare live vNN --col NAME' before publishing)" % ", ".join(odd[:8]))
    else:
        info.append("rows differ in number, so no cell-by-cell comparison")
    ok = not hard
    rep = ["# Check of v%d" % n, "", "Started from backup version %s (%s)." % (m["base"].get("backup_version"), m["base"].get("modified_by")),
           "", "Result: **%s**" % ("OK" if ok else "FAILED"), ""]
    if hard:
        rep += ["## Stops", ""] + ["- " + x for x in hard] + [""]
    if warn:
        rep += ["## To look at", ""] + ["- " + x for x in warn] + [""]
    rep += ["## Figures", ""] + ["- " + x for x in info] + [""]
    if changes:
        rep += ["## Cells changed per column", ""] + ["- %s: %d" % kv for kv in sorted(changes.items(), key=lambda x: -x[1])] + [""]
    (rd / "check_report.md").write_text("\n".join(rep), encoding="utf-8")
    m["check"] = {"ok": ok, "at": now(), "stops": hard, "warnings": warn, "info": info}
    save_manifest(rd, m)
    print()
    for x in info:
        print("  " + x)
    for x in warn:
        print("  WARNING: " + x)
    for x in hard:
        print("  STOP   : " + x)
    print("\n  Check %s  (report: %s)\n" % ("OK" if ok else "FAILED", rd / "check_report.md"))
    if not ok:
        sys.exit(3)
    log("Next: python admin.py excel")


# ----------------------------------------------------------------------------------------------- excel
PS_SAVE = r'''
# Opens the workbook in Excel, lets it calculate, saves and quits. Every Excel call is retried: right after a big
# open, Excel often answers "busy" (0x800AC472) for a few seconds.
$path = $args[0]
$visible = ($args[1] -eq "1")
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $visible
$xl.DisplayAlerts = $false
$xl.EnableEvents = $false
$wb = $null
$code = 0
function Step([string]$name, [scriptblock]$sb, [int]$tries = 12) {
  for ($i = 1; $i -le $tries; $i++) {
    try { $null = (& $sb); Write-Host ("STEP ok: " + $name); return $true }
    catch { Start-Sleep -Seconds 5 }
  }
  Write-Host ("STEP FAILED: " + $name)
  return $false
}
Step "ready" { if (-not $xl.Ready) { throw "not ready" } } 24 | Out-Null
if (-not (Step "open" { $global:wb = $xl.Workbooks.Open($path, 0, $false) })) { $code = 1 }
if ($code -eq 0) {
  Step "automatic calculation" { $xl.Calculation = -4105 } | Out-Null
  Step "calculate" { $xl.Calculate() } | Out-Null
  for ($i = 0; $i -lt 120; $i++) {                      # wait until Excel has finished calculating (up to 10 min)
    try { if ($xl.CalculationState -eq 0) { break } } catch {}
    Start-Sleep -Seconds 5
  }
  if (-not (Step "save" { $global:wb.Saved = $false; $global:wb.Save() })) { $code = 1 }
  Step "close" { $global:wb.Close($false) } | Out-Null
}
Step "quit" { $xl.Quit() } 6 | Out-Null
try { [System.Runtime.Interopservices.Marshal]::ReleaseComObject($xl) | Out-Null } catch {}
[GC]::Collect(); [GC]::WaitForPendingFinalizers()
exit $code
'''


def store_check_dp_values(path):
    """Store the value of every check_dp formula next to the formula, without Excel.

    The builder writes check_dp as a live formula (Country & Province & District & SIA_date & total & Response) with no
    stored result. Excel normally stores it when the file is saved once. This does the same thing directly in the file:
    the formula stays, its result is written beside it (dates as Excel's serial number, numbers in General format),
    so every reader (Excel, R, pandas) sees the value. Returns the number of cells filled."""
    import zipfile
    from xml.sax.saxutils import escape, unescape
    path = Path(path)
    tmp = path.with_name(path.stem + ".tmp_values.xlsx")
    with zipfile.ZipFile(path) as zin:
        sheets = sorted(n for n in zin.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml$", n))
        if not sheets:
            raise ValueError("no worksheet found in the file")
        name = sheets[0]
        xml = zin.read(name).decode("utf-8")
        if "<sharedStrings" in xml or 't="s"' in xml:
            raise ValueError("this file uses shared strings: open it in Excel and save it instead")
        cell_re = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
        row_re = re.compile(r"(<row [^>]*>)(.*?)(</row>)", re.S)
        n_done = 0

        def text_of(attrs, inner):
            if not inner:
                return ""
            if "inlineStr" in attrs:
                m = re.search(r"<t[^>]*>(.*?)</t>", inner, re.S)
                return unescape(m.group(1)) if m else ""
            m = re.search(r"<v>(.*?)</v>", inner, re.S)
            if not m or m.group(1) == "":
                return ""
            v = unescape(m.group(1))
            if 't="b"' in attrs:
                return "TRUE" if v in ("1", "true") else "FALSE"
            if 't="str"' in attrs or 't="e"' in attrs:
                return v
            try:
                x = float(v)
                return str(int(x)) if x == int(x) and abs(x) < 1e15 else ("%.15g" % x)
            except ValueError:
                return v

        def do_row(m):
            nonlocal n_done
            head, body, tail = m.group(1), m.group(2), m.group(3)
            if "<f>" not in body:
                return m.group(0)
            cells = {}
            for c in cell_re.finditer(body):
                cells[c.group(1)] = text_of(c.group(3), c.group(4))

            def fill(c):
                nonlocal n_done
                inner = c.group(4) or ""
                f = re.match(r"<f>([^<]*)</f><v\s*/?>(?:</v>)?$", inner)
                if not f:
                    return c.group(0)
                refs = re.findall(r"([A-Z]+)\d+", unescape(f.group(1)))
                val = "".join(cells.get(r, "") for r in refs)
                attrs = re.sub(r'\s+t="[^"]*"', "", c.group(3))
                n_done += 1
                return '<c r="%s%s"%s t="str"><f>%s</f><v>%s</v></c>' % (c.group(1), c.group(2), attrs, f.group(1), escape(val))
            return head + cell_re.sub(fill, body) + tail
        new_xml = row_re.sub(do_row, xml)
        if n_done == 0:
            raise ValueError("no formula without a stored value was found")
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, new_xml.encode("utf-8") if item.filename == name else zin.read(item.filename))
    os.replace(tmp, path)
    return n_done


def checkdp_filled(path):
    """True when the check_dp column holds stored values (not empty) at the top, the middle and the end."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        head = next(it)
        if "check_dp" not in head:
            return False, "no check_dp column"
        j = head.index("check_dp")
        vals = [r[j] if j < len(r) else None for r in it]
    finally:
        wb.close()
    if not vals:
        return False, "no data rows"
    probe = vals[:200] + vals[len(vals) // 2: len(vals) // 2 + 200] + vals[-200:]
    empty = sum(1 for v in probe if v in (None, ""))
    return empty == 0, "%d of %d probed cells empty" % (empty, len(probe))


def cmd_excel(a):
    """Store the check_dp values. Default: written straight into the file, formula kept (no Excel needed).
    --use-excel: desktop Excel opens, calculates and saves it (falls back to the direct method if Excel fails).
    --verify-only: you saved the file yourself; only check the values."""
    n, rd, m = latest_run(unpublished=True)
    if not n:
        die("no build waiting: run 'build' first.")
    if not m.get("check", {}).get("ok"):
        die("v%d has not passed 'check' yet." % n)
    f = Path(m["build"]["file"])
    method = "manual"

    def verify():
        try:
            return checkdp_filled(f)
        except Exception as e:
            return False, "the file could not be read (%s): close any Excel window that still has it open" % type(e).__name__

    def python_store():
        log("Storing the check_dp values in the file (the formulas stay) ...")
        try:
            k = store_check_dp_values(f)
        except Exception as e:
            die("could not store the values: %s. Open the file in Excel, save once and run 'excel --verify-only'." % e)
        log("check_dp: %d values stored." % k)

    if not a.verify_only:
        want_excel = a.use_excel or str(S.get("excel_method", "python")).lower() == "excel"
        if not want_excel:
            ok, why = verify()
            if not ok:
                python_store()
            method = "values stored by python (formulas kept)"
        else:
            if os.name != "nt":
                die("the Excel method needs Windows with Excel.")
            ps = rd / "_save_in_excel.ps1"
            ps.write_text(PS_SAVE, encoding="utf-8")
            log("Opening %s in Excel, recalculating and saving (a minute or two) ..." % f.name)
            try:
                r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps), str(f),
                                    "1" if S.get("excel_visible", True) else "0"],
                                   capture_output=True, text=True, errors="replace",
                                   timeout=int(S["excel_timeout_min"]) * 60)
                if r.returncode:
                    log("Excel reported a problem: %s" % (r.stdout.strip() or r.stderr.strip())[-300:].replace("\n", " | "))
            except subprocess.TimeoutExpired:
                log("Excel did not finish in %s minutes." % S["excel_timeout_min"])
            try:
                ps.unlink()
            except Exception:
                pass
            ok, why = verify()
            method = "excel automation"
            if not ok:
                log("Excel did not store the values (a copy of Office that is not activated opens files read-only). "
                    "Storing them directly instead.")
                python_store()
                method = "values stored by python (formulas kept) after Excel failed"
    ok, why = verify()
    m["excel"] = {"ok": ok, "at": now(), "method": method, "detail": why}
    save_manifest(rd, m)
    if not ok:
        die("the check_dp values are still empty (%s). Open the file in Excel, save it, and run 'excel --verify-only'." % why)
    log("check_dp OK: %s (%s)." % (why, method))
    log("Next: python admin.py publish")


# ----------------------------------------------------------------------------------------------- publish
def cmd_publish(a):
    n, rd, m = latest_run(unpublished=True)
    if not n:
        die("no build waiting to be published.")
    f = Path(m["build"]["file"])
    if not m.get("check", {}).get("ok"):
        die("v%d has not passed 'check' (run it, and fix what it stops on)." % n)
    if not m.get("excel", {}).get("ok"):
        die("v%d was not saved in Excel yet (run 'excel')." % n)
    if not f.is_file():
        die("the build file is missing: %s" % f)
    ok, why = checkdp_filled(f)
    if not ok:
        die("the check_dp values of %s are empty (%s). Run 'excel' again." % (f.name, why))
    # the live file must still be the one the build started from
    row, snap = do_backup(show=False)
    log("Comparing the live file with the one v%d started from ..." % n)
    live = load_norm(snap)
    bp = Path(m["base"]["path"])
    base_fp = fingerprint(load_norm(bp)) if bp.is_file() else m["base"]["fingerprint"]   # same normalisation on both sides
    same = fingerprint(live) == base_fp
    if not same:
        base_cols = m["base"]["columns"]
        extra = [c for c in live.columns if c not in base_cols]
        die("the live SharePoint file changed after v%d was built (now backup version %s, saved by %s). "
            "Publishing now would overwrite that change. Run 'build' again: it starts from the current live file.%s"
            % (n, row["version"], row["modified_by"], (" New columns: %s." % ", ".join(extra)) if extra else ""))
    log("The live file is the one v%d started from (backup version %s)." % (n, row["version"]))
    pub = P(S["publisher"])
    cmd = [sys.executable, pub, "--file", f, "--backup-repo", BACKUP_REPO] + list(a.extra or [])
    rc, _ = run_cmd(cmd, cwd=HERE, tee=rd / "publish.log")
    if rc:
        die("the publish script stopped (exit %s). Nothing is marked as published; see %s." % (rc, rd / "publish.log"))
    row2, _ = latest_snapshot()
    m["published"] = {"at": now(), "backup_version": row2["version"] if row2 else None, "file": str(f)}
    save_manifest(rd, m)
    st = read_json(STATE, {})
    st["published"] = {"version": n, "at": m["published"]["at"], "backup_version": m["published"]["backup_version"]}
    write_json(STATE, st)
    log("v%d is live. Local backup of the new live file: version %s." % (n, m["published"]["backup_version"]))
    log("Next: commit the code if it changed ('commit'), and tell the team the new version is live.")


# ----------------------------------------------------------------------------------------------- all
def cmd_all(a):
    class X:  # reuse the sub-commands with their defaults
        pass
    b = X(); b.base = None; b.no_backup = False; b.extra = a.extra
    cmd_build(b)
    c = X(); c.run = None
    cmd_check(c)
    e = X(); e.verify_only = False
    cmd_excel(e)
    p = X(); p.extra = []
    cmd_publish(p)


# ----------------------------------------------------------------------------------------------- git
def cmd_commit(a):
    if not a.files:
        die("name the files to commit, e.g.: python admin.py commit -m \"message\" add_coadmin_columns.py README.md")
    low = a.message.lower()
    if any(x in low for x in FORBIDDEN_MSG):
        die("the message must not contain Co-Authored-By / Claude-Session lines.")
    for f in a.files:
        name = Path(f).name.lower()
        if name.endswith(DATA_EXT) and name not in DATA_OK:
            die("%s is a data file: data is never committed (only lookup.xlsx and geo_reference.csv)." % f)
        if "secrets" in name:
            die("%s looks like a secrets file: never committed." % f)
    rc, out = git("add", "--", *a.files)
    if rc:
        die("git add failed: %s" % out)
    rc, out = git("diff", "--cached", "--quiet")
    if rc == 0:
        die("nothing to commit for these files (no change).")
    rc, out = git("commit", "-m", a.message)
    print(out)
    if rc:
        die("git commit failed.")
    if not a.no_push:
        rc, out = git("push")
        print(out)
        if rc:
            die("the commit is made but the push failed: run 'git push' yourself.")
    log("Committed%s: %s" % ("" if a.no_push else " and pushed", a.message))


# ----------------------------------------------------------------------------------------------- scope sync, doctor
def cmd_sync_scope(a):
    sd, dd = S.get("scope_script_dir"), S.get("scope_data_dir")
    if not (sd or dd):
        die("set scope_script_dir / scope_data_dir in admin_settings.json first.")
    if sd:
        Path(sd).mkdir(parents=True, exist_ok=True)
        for name in (S["builder"], S["publisher"]):
            shutil.copy2(P(name), Path(sd) / Path(name).name)
            log("copied %s -> %s" % (name, sd))
    n, d, m = latest_run()
    if dd and n:
        up = Path(dd) / "upload_to_sharepoint"
        up.mkdir(parents=True, exist_ok=True)
        for p in d.glob("*_v%d.xlsx" % n):
            tgt = up if ("AFRO_admin_data_v" in p.name or "dictionary" in p.name) else Path(dd)
            shutil.copy2(p, tgt / p.name)
        log("copied the v%d files to %s" % (n, dd))


def cmd_doctor(a):
    ok = True

    def line(good, msg):
        nonlocal ok
        ok &= good
        print("  [%s] %s" % ("ok" if good else "!!", msg))
    print("\nChecking the set-up:\n")
    line(sys.version_info >= (3, 9), "python %s" % sys.version.split()[0])
    for mod in ("pandas", "openpyxl", "numpy", "sklearn"):
        try:
            __import__(mod)
            line(True, "package %s" % mod)
        except Exception:
            line(mod == "sklearn", "package %s%s" % (mod, " (optional)" if mod == "sklearn" else " - install with: pip install " + mod))
    for label, p in (("builder", P(S["builder"])), ("publisher", P(S["publisher"])), ("lookup", P(S["lookup"])),
                     ("geo reference", P(S["geo"])), ("backup tool", BACKUP_REPO / "sp_versions.py")):
        line(p.is_file(), "%s: %s" % (label, p))
    line((BACKUP_REPO / "config" / "secrets.env").is_file(), "SharePoint credentials file present (not read here)")
    line(shutil.which("git") is not None, "git")
    if os.name == "nt":
        try:
            r = subprocess.run(["powershell", "-NoProfile", "-Command",
                                "[bool](Get-ItemProperty 'HKLM:\\SOFTWARE\\Classes\\Excel.Application' -ErrorAction SilentlyContinue)"],
                               capture_output=True, text=True, timeout=30)
            line("True" in r.stdout, "Excel installed (for the automatic save)")
        except Exception:
            line(False, "Excel could not be detected")
    else:
        print("  [--] Excel step is for Windows only")
    print("\n  %s\n" % ("All good." if ok else "Fix the lines marked !! first."))


# ----------------------------------------------------------------------------------------------- sources
SHOW_COLS = ["Country", "Province", "District", "SIA_date", "Round_Add", "Vaccine_type", "Response", "PopPolio",
             "TotalNbrVaccPolio", "CVPolio"]


def resolve_source(spec):
    """'live' (newest local backup), 'vNN' (a build), 'bN' (backup version N) or a file path."""
    if not spec or spec == "live":
        row, path = latest_snapshot()
        if row is None:
            die("no local backup yet: run 'backup'.")
        return path, "backup version %s (%s)" % (row["version"], row["modified_by"])
    m = re.match(r"^v(\d+)$", spec)
    if m:
        d = RUNS / ("v%s" % m.group(1))
        mf = read_json(d / "manifest.json", {})
        if not mf.get("build", {}).get("file"):
            die("build %s not found." % spec)
        return Path(mf["build"]["file"]), "build %s" % spec
    m = re.match(r"^b(\d+)$", spec)
    if m:
        for r in catalog_rows("admin"):
            if r["version"] == m.group(1):
                return BACKUP_REPO / r["snapshot"], "backup version %s (%s)" % (r["version"], r["modified_by"])
        die("backup version %s not found." % m.group(1))
    p = Path(spec)
    if not p.is_file():
        die("file not found: %s" % spec)
    return p, p.name


# ----------------------------------------------------------------------------------------------- inspect
def hygiene(df, builder, lk):
    """Checks that are not in the portal quality report: stray characters, names, the lookup."""
    out = []
    for c in ("Country", "Province", "District", "Response", "Vaccine_type", "Round_Add"):
        if c in df.columns:
            bad = df[c].astype(str).str.contains(r"_x000D_|\s\s|^\s|\s$|[\r\n\t]", regex=True, na=False)
            if bad.any():
                out.append(("Stray characters or spaces in %s" % c, int(bad.sum()), ", ".join(sorted(set(df.loc[bad, c].astype(str).str.strip()))[:4])))
    if "Response" in df.columns:
        bad = response_country_mismatches(df, builder, lk)
        if bad:
            out.append(("Response of another country (per the lookup)", sum(bad.values()),
                        "; ".join("%s / %s" % k for k in list(bad)[:4])))
    for c in ("Country", "Province", "District", "SIA_date", "Round_Add", "Vaccine_type", "Response"):
        if c in df.columns:
            e = int((df[c].astype(str).str.strip() == "").sum())
            if e:
                out.append(("Empty %s" % c, e, ""))
    return out


def cmd_inspect(a):
    path, label = resolve_source(a.source)
    log("Inspecting %s" % label)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    od = WORK / "inspect" / stamp
    od.mkdir(parents=True, exist_ok=True)
    cmd = [sys.executable, P(S["builder"]), path, "--check-only", "--out-dir", od, "--lookup", P(S["lookup"])]
    if P(S["geo"]).is_file():
        cmd += ["--geo", P(S["geo"])]
    if a.country:
        cmd += ["--country", a.country]
    rc, text = run_cmd(cmd, cwd=HERE, tee=od / "inspect.log")
    if rc:
        die("the quality check failed (see %s)." % (od / "inspect.log"))
    pd = _pd()
    builder = import_builder()
    lk = builder.load_lookup(str(P(S["lookup"])))
    df = load_norm(path)
    if a.country:
        df = df[df["Country"].map(builder._qcountry) == builder._qcountry(a.country)]
    print("\nHygiene checks:")
    hy = hygiene(df, builder, lk)
    for t, k, ex in hy:
        print("  %-50s %6d  %s" % (t, k, ex))
    if not hy:
        print("  nothing found")
    rep, nodata = builder.lookup_crosscheck(df, lk)
    bad = rep[rep["Status"] != "OK"]
    print("\nCampaign rounds against the lookup: %d of %d agree" % (len(rep) - len(bad), len(rep)))
    for st, k in bad["Status"].value_counts().items():
        print("  %-55s %4d rounds" % (st, k))
    with pd.ExcelWriter(od / "lookup_check.xlsx") as w:
        rep.to_excel(w, sheet_name="Lookup_check", index=False)
        nodata.to_excel(w, sheet_name="Lookup_rounds_without_data", index=False)
    print("\nReports in %s" % od)
    print("  AFRO_admin_data_quality_report.xlsx (portal checks), lookup_check.xlsx, inspect.log\n")


# ----------------------------------------------------------------------------------------------- find
def cmd_find(a):
    path, label = resolve_source(a.source)
    df = load_norm(path)
    df.index = df.index + 2          # the Excel row number
    m = _pd().Series(True, index=df.index)
    for col, val in (("Country", a.country), ("Province", a.province), ("District", a.district),
                     ("Response", a.response), ("Vaccine_type", a.vaccine)):
        if val and col in df.columns:
            m &= df[col].str.upper().str.contains(re.escape(val.upper()), regex=True, na=False)
    if a.round and "Round_Add" in df.columns:
        m &= df["Round_Add"].str.upper() == a.round.upper()
    if a.date and "SIA_date" in df.columns:
        m &= df["SIA_date"].str.startswith(a.date)
    r = df[m]
    cols = [c for c in SHOW_COLS if c in r.columns] + [c for c in (a.cols or []) if c in r.columns]
    print("\n%s: %d record(s) match\n" % (label, len(r)))
    if r.empty:
        return
    if a.summary:
        g = r.assign(_t=_pd().to_numeric(r.get("TotalNbrVaccPolio"), errors="coerce")) \
             .groupby([c for c in ("Country", "Response", "Round_Add", "SIA_date", "Vaccine_type") if c in r.columns], dropna=False) \
             .agg(records=("Country", "size"), vaccinated=("_t", "sum")).reset_index()
        g["SIA_date"] = g["SIA_date"].str[:10]
        print(g.to_string(index=False))
    else:
        t = r[cols].head(a.limit).copy()
        if "SIA_date" in t.columns:
            t["SIA_date"] = t["SIA_date"].str[:10]
        print(t.to_string())
        if len(r) > a.limit:
            print("\n... %d more (use --limit N, or --summary)" % (len(r) - a.limit))
    print()


# ----------------------------------------------------------------------------------------------- compare
def cmd_compare(a):
    pa, la = resolve_source(a.a)
    pb, lb = resolve_source(a.b)
    log("Comparing %s with %s" % (la, lb))
    pd = _pd()
    A, B = load_norm(pa), load_norm(pb)
    print("\n  A = %s: %d records, %d columns" % (la, len(A), len(A.columns)))
    print("  B = %s: %d records, %d columns" % (lb, len(B), len(B.columns)))
    only_a = [c for c in A.columns if c not in B.columns]
    only_b = [c for c in B.columns if c not in A.columns]
    if only_a:
        print("  columns only in A: %s" % ", ".join(only_a[:12]))
    if only_b:
        print("  columns only in B: %s" % ", ".join(only_b[:12]))
    common = [c for c in A.columns if c in B.columns]
    keys = [k for k in S["key_columns"] if k in common]
    out = WORK / "compare"
    out.mkdir(parents=True, exist_ok=True)
    name = out / ("compare_%s_vs_%s.xlsx" % (re.sub(r"\W+", "", a.a or "live"), re.sub(r"\W+", "", a.b or "live")))
    sheets = {}
    if len(A) == len(B):
        diffs = {}
        for c in common:
            ne = A[c].values != B[c].values
            if ne.any():
                diffs[c] = ne
        print("\n  same number of records: cells that differ, by column")
        if not diffs:
            print("    none: the two versions hold the same data")
        rows = []
        for c, ne in sorted(diffs.items(), key=lambda kv: -kv[1].sum()):
            print("    %-28s %6d" % (c, ne.sum()))
            if a.col is None or a.col == c:
                idx = ne.nonzero()[0][: a.limit if a.col else 200]
                for i in idx:
                    rows.append({"Excel_row": i + 2, "Column": c, "A": A[c].values[i], "B": B[c].values[i],
                                 **{k: A[k].values[i] for k in keys}})
        if rows:
            sheets["Changed_cells"] = pd.DataFrame(rows)
    else:
        def keyed(D):
            k = D[keys].agg("|".join, axis=1)
            return k + "|#" + k.groupby(k).cumcount().astype(str)
        ka, kb = keyed(A), keyed(B)
        ra, rb = A[ka.isin(set(kb)) == False], B[kb.isin(set(ka)) == False]
        print("\n  records only in A: %d    only in B: %d    in both: %d" % (len(ra), len(rb), len(A) - len(ra)))
        sheets["Only_in_A"] = ra[keys + [c for c in ("PopPolio", "TotalNbrVaccPolio") if c in ra.columns]].head(5000)
        sheets["Only_in_B"] = rb[keys + [c for c in ("PopPolio", "TotalNbrVaccPolio") if c in rb.columns]].head(5000)
        for lab, D in (("A", ra), ("B", rb)):
            if len(D):
                print("    only in %s, by Country / Response:" % lab)
                print("      " + D.groupby(["Country", "Response"]).size().sort_values(ascending=False).head(8).to_string().replace("\n", "\n      "))
    if sheets:
        with pd.ExcelWriter(name) as w:
            for k, v in sheets.items():
                v.to_excel(w, sheet_name=k, index=False)
        print("\n  details: %s\n" % name)


# ----------------------------------------------------------------------------------------------- rules / knowledge
RULES = [
    # (id, area, what the system does, where it comes from / why, switch)
    ("coadmin", "Co-administration", "Same district, round and Response with an nOPV2 record and a bOPV record are merged into one 'nOPV2 & bOPV' record; TotalNbrVaccPolio = the larger of the two (same children). A combined record gets the same total in both columns until the country gives the split.",
     "Feedback from Valentina Bigoni, WHO AFRO", "(always)"),
    ("unused-vaccine", "Co-administration", "A vaccine that was not used has an EMPTY total and coverage (not 0). Real reported zeros are kept.", "Decision 30 Oct 2026 (Ahmed's convention)", "(always)"),
    ("poppolio", "Population", "PopPolio is corrected from vaccinated / CVPolio when coverage is impossible (above 1.5, or below 0.5 with a plausible CVPolio); decimals rounded; otherwise the district median of the year; then a gradient-boosting estimate (listed for confirmation).", "Quality work with the focal points", "--no-fix-poppolio, --no-model-poppolio"),
    ("copy-repair", "Copy / drag-fill errors", "A Response of another country, or nOPV3..nOPV15, is restored in this order: the LOOKUP (country + vaccine + start date), the newest earlier backup, the other rows of the same country/date/round, the campaign footprint. Only unique answers are used; the rest is UNRESOLVED.", "The lookup is the primary reference (your rule)", "--no-repair, --no-infer"),
    ("copied-rounds", "Copy / drag-fill errors", "Rounds that the lookup does not have and that only copy another round of the same Response and date are removed and listed.", "Found in the 2026 data", "--keep-copied-rounds"),
    ("response-fill", "Key fields", "An empty Response is filled from the lookup (exact start date, or within 21 days with the same round number), only when unique.", "Portal check 'Missing key field'", "--no-fill-response"),
    ("response-names", "Key fields", "A Response that is not a lookup name but starts the name of the one lookup round of that country within 7 days is renamed to the lookup's spelling (MLW-2026-02-1_nOPV -> MLW-2026-02-1_nOPV-NID).", "Decision 30 Oct 2026: follow the lookup", "--no-align-response-names, --align-response-names COUNTRY"),
    ("cvpolio", "Coverage", "An empty CVPolio is filled with TotalNbrVaccPolio / PopPolio. Records with only a target (no vaccinated) stay empty.", "Request of 28 Oct 2026", "--no-fill-cvpolio"),
    ("district-fill", "Districts", "An empty District is filled from a recovered file, matched on country + Response + round + province: same PopPolio and total; same PopPolio and total within 10%; same total; total + the target implied by the coverage; the last unmatched pair; same PopPolio or total with unique pairing.", "Ethiopia file AFRO-ETH FIVE REGIONS R1 (282 of 282 filled)", "--district-source FILE"),
    ("addis", "Districts", "Addis Ababa rows that repeat the sub-city name (NIFAS SILK LAFTO x13) are renamed '<SUB-CITY> WOREDA NN' from the file, only when every row of the sub-city matches a different woreda. Rows without a source (round 2 and 4 sub-cities of round 1) are left; they stay possible duplicates until a woreda-level file exists.", "Request of 29 Oct 2026", "(with --district-source)"),
    ("province", "Geography", "A Province that belongs to another country (per geo_reference.csv) is corrected when certain.", "Admin portal checks", "--no-fix-province"),
    ("text-clean", "Hygiene", "Stray control characters (carriage returns), line breaks and double / edge spaces in Country, Province, District, Response, Vaccine_type and Round_Add are removed.", "Found in District (KAJIADO EAST) 30 Oct 2026", "--no-clean-text"),
    ("dates", "Hygiene", "SIA_date written as text ('18-sept-20') becomes a real date.", "Admin portal", "(always)"),
    ("duplicates", "Hygiene", "Exact copies are removed; records with the same key but different values are NOT removed, only reported (Duplicate record). Key = Country, Province, District, [Admin_3], SIA_date, Round_Add, Vaccine_type, Response.", "Admin portal", "--no-dedupe, --resolve-duplicates"),
    ("quality", "Checks", "The portal quality checks (missing key field, coverage above the ceiling, target jumps, report date, province/country, co-administration consistency, vaccine spelling, duplicates) produce a report with Errors / Warnings / Info.", "Admin portal (run_admin_quality_check.R)", "--no-quality"),
]
DECISIONS = [
    "The lookup file is the primary reference for campaign names, dates and rounds; backups come second.",
    "Never guess: an answer is used only when it is unique; otherwise the record is left and listed.",
    "A build always starts from the live SharePoint file, so nobody's newer edits are lost.",
    "SharePoint keeps no version history (the local backup has it); one published file, replaced in place.",
    "Data is never committed to git (only lookup.xlsx and geo_reference.csv). No attribution lines in commit messages.",
    "Each code change is its own commit, followed by its own push.",
    "Open questions kept visible: Addis Ababa round 2 and 4 sub-cities lack woreda names (need a woreda-level file); Tegedie exists in two zones of Amhara with the same name; about 1,460 Nigeria 2022-Jan 2023 records have no Response in the lookup.",
]


def rules_markdown():
    L = ["# What the admin data system knows", "", "Generated by `python admin.py rules --markdown`. The rules live in `add_coadmin_columns.py`; this is their catalogue.", ""]
    area = None
    for rid, ar, what, src, sw in RULES:
        if ar != area:
            L += ["", "## " + ar, ""]
            area = ar
        L += ["- **%s**: %s  \n  Source: %s. Switch: `%s`." % (rid, what, src, sw)]
    L += ["", "## Decisions and standing rules", ""] + ["- " + d for d in DECISIONS] + [""]
    return "\n".join(L)


def cmd_rules(a):
    if a.markdown:
        out = HERE / "KNOWLEDGE.md"
        out.write_text(rules_markdown(), encoding="utf-8")
        print("written: %s" % out)
        return
    print()
    for rid, ar, what, src, sw in RULES:
        if a.name and a.name.lower() not in (rid + " " + ar + " " + what).lower():
            continue
        print("[%s] %s\n    %s\n    source: %s | switch: %s\n" % (rid, ar, what, src, sw))
    if not a.name:
        print("Decisions and standing rules:")
        for d in DECISIONS:
            print("  - " + d)
        print()


# ----------------------------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd")
    s = sub.add_parser("status"); s.add_argument("--live", action="store_true", help="read SharePoint first (runs the backup)")
    sub.add_parser("runs")
    sub.add_parser("backup")
    s = sub.add_parser("build"); s.add_argument("--base", help="start from this file instead of the latest backup of the live file")
    s.add_argument("--no-backup", action="store_true", help="do not read SharePoint, use the newest local backup")
    s.add_argument("extra", nargs="*", help="extra options for the builder, after --")
    s = sub.add_parser("check"); s.add_argument("--run", type=int, help="version number (default: the newest)")
    s = sub.add_parser("excel"); s.add_argument("--verify-only", action="store_true", help="you saved the file in Excel yourself")
    s.add_argument("--use-excel", action="store_true", help="let desktop Excel store the values (default: written directly, no Excel)")
    s = sub.add_parser("publish"); s.add_argument("extra", nargs="*", help="extra options for the publish script, after --")
    s = sub.add_parser("all"); s.add_argument("extra", nargs="*")
    s = sub.add_parser("commit"); s.add_argument("-m", "--message", required=True); s.add_argument("--no-push", action="store_true")
    s.add_argument("files", nargs="*")
    sub.add_parser("sync-scope")
    sub.add_parser("doctor")
    s = sub.add_parser("inspect"); s.add_argument("--source", help="live (default), vNN build, bN backup version, or a file")
    s.add_argument("--country")
    s = sub.add_parser("find"); s.add_argument("--source"); s.add_argument("--country"); s.add_argument("--province")
    s.add_argument("--district"); s.add_argument("--response"); s.add_argument("--vaccine"); s.add_argument("--round")
    s.add_argument("--date", help="YYYY-MM-DD or YYYY-MM"); s.add_argument("--limit", type=int, default=40)
    s.add_argument("--summary", action="store_true", help="count per Response / round / date instead of listing records")
    s.add_argument("--cols", nargs="*", help="more columns to show")
    s = sub.add_parser("compare"); s.add_argument("a", help="live, vNN, bN or a file"); s.add_argument("b")
    s.add_argument("--col", help="list the changed cells of this column"); s.add_argument("--limit", type=int, default=200)
    s = sub.add_parser("rules"); s.add_argument("name", nargs="?"); s.add_argument("--markdown", action="store_true")
    a = ap.parse_args()
    if not a.cmd:
        ap.print_help()
        return
    fn = {"status": cmd_status, "runs": cmd_runs, "backup": cmd_backup, "build": cmd_build, "check": cmd_check,
          "excel": cmd_excel, "publish": cmd_publish, "all": cmd_all, "commit": cmd_commit,
          "sync-scope": cmd_sync_scope, "doctor": cmd_doctor, "inspect": cmd_inspect, "find": cmd_find,
          "compare": cmd_compare, "rules": cmd_rules}[a.cmd]
    log("admin.py %s" % " ".join(sys.argv[1:]), quiet=True)
    try:
        fn(a)
    except KeyboardInterrupt:
        die("interrupted.", 130)


if __name__ == "__main__":
    main()
