#!/usr/bin/env python3
"""
publish_admin_data.py - put the latest corrected admin data on SharePoint, replacing the file that is there.

What it does, in this order (nothing is written to SharePoint until every check has passed and you confirmed):

  1. Finds the latest corrected version (AFRO_admin_data_v18.xlsx > v17 ...) in the output folder of
     add_coadmin_columns.py, or the file given with --file.
  2. Checks it: one sheet "Sheet1", the key columns and the co-administration columns, the check_dp values stored
     (the file was opened and saved once in Excel), and that it is not a country file.
  3. Downloads the file that is on SharePoint now and compares: same columns, row count not dropping by more than
     10%, not identical to what is already there.
  4. Backs up the live file locally with the SharePoint_version_backup tool and checks that the backup really holds
     it (SHA-256). If the backup is missing the script stops.
  5. Uploads the new file under the name AFRO_admin_data.xlsx INTO THE EXISTING SharePoint file (same file, same
     link, same place). The upload is refused if somebody changed the file in the meantime.
  6. Checks the upload (size and SHA-256 of what SharePoint now holds).
  7. Storage: deletes every historical version SharePoint kept of that file (SharePoint stores each overwrite as a
     full extra copy; the local backup keeps the history instead) and, when permitted, empties the deleted copies
     from the site recycle bin so the space is really freed.
  8. Runs the local backup again so the new live file is recorded as the newest local version.
  9. The data dictionary: the AFRO_admin_data_dictionary.xlsx that goes with the file is checked and pushed INTO THE
     EXISTING SharePoint dictionary the same way (local copy of the old one kept, upload verified, version history
     removed). A problem with the dictionary never undoes the admin data that is already published; the script
     then ends with exit code 3 and tells you so. --no-dictionary skips it, --dictionary-only pushes only that.

Usage (from this folder):
    python publish_admin_data.py                 # latest version, asks for confirmation
    python publish_admin_data.py --dry-run       # all checks, nothing written anywhere
    python publish_admin_data.py --file path\\to\\AFRO_admin_data_v18.xlsx
    python publish_admin_data.py --yes           # no question (all checks still apply)
    python publish_admin_data.py --dictionary-only --dictionary AFRO_admin_data_dictionary_v24.xlsx

Credentials: SHAREPOINT_TENANT_ID / SHAREPOINT_CLIENT_ID / SHAREPOINT_CLIENT_SECRET, taken from the environment or
from a secrets.env file (--secrets FILE, config/secrets.env next to this script, or the SharePoint_version_backup /
IM workflow config folders). The values are never printed.

Needs Python 3, requests and openpyxl.
"""
import argparse
import base64
import csv
import datetime as dt
import hashlib
import itertools
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

try:
    import requests
except ImportError:  # pragma: no cover
    sys.exit("The 'requests' package is missing:  pip install requests")
try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("The 'openpyxl' package is missing:  pip install openpyxl")

HERE = Path(__file__).resolve().parent

# ----------------------------------------------------------------------------------------------- settings
TARGET_NAME = "AFRO_admin_data.xlsx"
# The live file is found through its share link (the same one the backup tool uses), so its location can change.
ADMIN_SHARE_URL = ("https://worldhealthorg.sharepoint.com/:x:/r/sites/AF-pep/GISWORKSPACE/_layouts/15/Doc.aspx"
                   "?sourcedoc=%7BF9F9983C-76E8-4133-BF13-F69F9F320C15%7D&file=AFRO_admin_data.xlsx"
                   "&action=default&mobileredirect=true")
BACKUP_KEY = "admin"
DICT_NAME = "AFRO_admin_data_dictionary.xlsx"
DICT_SHARE_URL = ("https://worldhealthorg.sharepoint.com/:x:/r/sites/AF-pep/GISWORKSPACE/_layouts/15/Doc.aspx"
                  "?sourcedoc=%7B58972B92-BECE-4456-B465-DDFB95CCD411%7D&file=AFRO_admin_data_dictionary.xlsx"
                  "&action=default&mobileredirect=true&wdwpf=doclib-t")
DICT_SHEET = "Data dictionary"
EXIT_DICT_FAILED = 3          # the admin data is published but the dictionary push failed

# Where add_coadmin_columns.py leaves the upload files (the first existing folders are searched).
DEFAULT_DIRS = [
    HERE / "coadmin_output" / "upload_to_sharepoint",
    Path(r"C:\Users\TOURE\Documents\REPOSITORIES\input\scope\data\co_administration\upload_to_sharepoint"),
]
DEFAULT_BACKUP_REPO = Path(r"C:\Users\TOURE\Documents\Gith_repositories\SharePoint_version_backup")
SECRET_FILES = [
    HERE / "config" / "secrets.env",
    DEFAULT_BACKUP_REPO / "config" / "secrets.env",
    Path(r"C:\Users\TOURE\Documents\Gith_repositories\AFRO-Independant-Monitoring-data-workflow\config\secrets.env"),
]

REQUIRED_COLUMNS = ["Country", "Province", "District", "SIA_date", "Response", "Round_Add", "Year", "PopPolio",
                    "TotalNbrVaccPolio", "CVPolio", "Vaccine_type",
                    "Tot_Vacc_nOPV2", "CV_nOPV2", "Total_Vacc_bOPV", "CV_bOPV", "check_dp"]
MIN_COUNTRIES = 25            # fewer countries = a country file, not the regional file
MAX_ROW_DROP = 0.10           # refuse when the new file has more than 10% fewer rows than the live one
CHUNK = 320 * 1024 * 32       # 10 MiB, a multiple of 320 KiB as Graph requires
GRAPH = os.environ.get("PUBLISH_GRAPH_BASE", "https://graph.microsoft.com/v1.0")      # (test hooks)
GRAPH_BETA = os.environ.get("PUBLISH_GRAPH_BETA", "https://graph.microsoft.com/beta")
TOKEN_URL = os.environ.get("PUBLISH_TOKEN_URL", "https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token")

LOGDIR = HERE / "logs"


# ----------------------------------------------------------------------------------------------- helpers
def log(msg=""):
    line = f"{dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  {msg}" if msg else ""
    print(line, flush=True)
    try:
        LOGDIR.mkdir(exist_ok=True)
        with open(LOGDIR / f"publish_{dt.date.today():%Y%m%d}.log", "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except OSError:
        pass


class SoftStop(Exception):
    """Raised instead of exiting while the dictionary is handled: the admin data is already published then."""


def die(msg, code=1):
    if STATE.get("soft"):
        raise SoftStop(msg)
    log("")
    log("STOPPED - " + msg)
    log("Nothing was changed on SharePoint." if not STATE["uploaded"] else "The upload had already happened, see above.")
    sys.exit(code)


STATE = {"uploaded": False}


def human(n):
    n = float(n)
    for u in ("B", "KB", "MB", "GB"):
        if n < 1024 or u == "GB":
            return f"{n:.0f} {u}" if u == "B" else f"{n:.1f} {u}"
        n /= 1024


def sha256_file(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def load_secrets(extra=None):
    """Fill the SHAREPOINT_* variables from the first secrets.env that has them. Values are never printed."""
    need = ("SHAREPOINT_TENANT_ID", "SHAREPOINT_CLIENT_ID", "SHAREPOINT_CLIENT_SECRET")
    if all(os.environ.get(k) for k in need):
        return "environment"
    for p in ([Path(extra)] if extra else []) + SECRET_FILES:
        try:
            if not p.exists():
                continue
            for line in p.read_text(encoding="utf-8").splitlines():
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, _, v = line.partition("=")
                    if k.strip() in need and v.strip():
                        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))
            if all(os.environ.get(k) for k in need):
                return str(p)
        except OSError:
            continue
    return None


# ----------------------------------------------------------------------------------------------- 1. which file
VERSIONED = re.compile(r"^AFRO_admin_data_v(\d+)\.xlsx$", re.I)
PLAIN = re.compile(r"^AFRO_admin_data\.xlsx$", re.I)


def find_candidates(dirs):
    found = []
    for d in dirs:
        d = Path(d)
        if not d.is_dir():
            continue
        for p in d.iterdir():
            m = VERSIONED.match(p.name)
            if m:
                found.append((int(m.group(1)), p.stat().st_mtime, p))
            elif PLAIN.match(p.name):
                found.append((None, p.stat().st_mtime, p))
    return found


def choose_file(args):
    if args.file:
        p = Path(args.file)
        if not p.is_file():
            die(f"the file {p} does not exist.")
        return p, []
    dirs = [Path(args.dir)] if args.dir else DEFAULT_DIRS
    cand = find_candidates(dirs)
    if not cand:
        die("no AFRO_admin_data_vNN.xlsx / AFRO_admin_data.xlsx found in:\n    " + "\n    ".join(map(str, dirs))
            + "\n  Give the file with --file or the folder with --dir.")
    versioned = [c for c in cand if c[0] is not None]
    pool = versioned or cand
    pool.sort(key=lambda c: ((c[0] if c[0] is not None else -1), c[1]))
    return pool[-1][2], sorted(cand, key=lambda c: ((c[0] if c[0] is not None else -1), c[1]), reverse=True)


# ----------------------------------------------------------------------------------------------- 2. check a file
def inspect_workbook(path, sample=300):
    """Header, row count, number of sheets, and a look at the first rows (country count is read from the Country column
    of the whole file only when cheap). Uses read-only mode, so an 11 MB file takes a few seconds."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        ws = wb[names[0]]
        rows = ws.iter_rows(values_only=True)
        header = [("" if v is None else str(v)) for v in next(rows)]
        n_rows = ws.max_row - 1 if ws.max_row else None
        first = list(itertools.islice(rows, sample))
        pos = {h: i for i, h in enumerate(header)}
        cd_vals = [r[pos["check_dp"]] for r in first if "check_dp" in pos and len(r) > pos["check_dp"]]
        cd_filled = sum(v not in (None, "") for v in cd_vals)
        info = {"sheets": names, "header": header, "rows": n_rows, "sample": len(first),
                "check_dp_filled": cd_filled, "check_dp_sample": len(cd_vals)}
        # countries: scan the Country column of the whole sheet (one pass, only that column)
        if "Country" in pos:
            ci = pos["Country"] + 1
            ctry, last_non_empty = set(), 0
            n = 0
            for (v,) in ws.iter_rows(min_row=2, min_col=ci, max_col=ci, values_only=True):
                n += 1
                if v not in (None, ""):
                    ctry.add(str(v).strip().upper())
                    last_non_empty = n
            info["countries"] = len(ctry)
            info["rows"] = last_non_empty        # trailing blank rows are not records
        return info
    finally:
        wb.close()


def check_new_file(path, info):
    problems, warnings = [], []
    if info["sheets"] != ["Sheet1"]:
        warnings.append(f"sheets are {info['sheets']} (expected only Sheet1)")
    missing = [c for c in REQUIRED_COLUMNS if c not in info["header"]]
    if missing:
        problems.append("columns missing: " + ", ".join(missing))
    dup = sorted({h for h in info["header"] if h and info["header"].count(h) > 1})
    if dup:
        problems.append("duplicated column names: " + ", ".join(dup))
    if info.get("countries") is not None and info["countries"] < MIN_COUNTRIES:
        problems.append(f"only {info['countries']} countries - this looks like a country file, not the regional file")
    if not info.get("rows"):
        problems.append("no data rows")
    if "check_dp" in info["header"] and info["check_dp_sample"] and info["check_dp_filled"] == 0:
        problems.append("the check_dp values are empty: open the file in Excel, save once (Ctrl+S) and run again "
                        "(use --allow-unsaved to skip this check)")
    return problems, warnings


# ----------------------------------------------------------------------------------------------- SharePoint access
class Graph:
    def __init__(self, secrets_file=None):
        src = load_secrets(secrets_file)
        if not src:
            die("the SharePoint credentials were not found. Set SHAREPOINT_TENANT_ID / SHAREPOINT_CLIENT_ID / "
                "SHAREPOINT_CLIENT_SECRET, or give a secrets.env with --secrets FILE.")
        self.secrets_source = src
        self._token_at = 0.0
        self._token = None
        self.refresh()

    def refresh(self):
        t, c, s = (os.environ[k].strip() for k in ("SHAREPOINT_TENANT_ID", "SHAREPOINT_CLIENT_ID",
                                                   "SHAREPOINT_CLIENT_SECRET"))
        try:
            r = requests.post(TOKEN_URL.format(tenant=t),
                              data={"client_id": c, "client_secret": s, "grant_type": "client_credentials",
                                    "scope": "https://graph.microsoft.com/.default"}, timeout=30)
            r.raise_for_status()
            self._token = r.json()["access_token"]
            self._token_at = time.time()
        except Exception as e:
            die(f"could not sign in to SharePoint ({type(e).__name__}). Check the credentials and the network.")

    def headers(self, extra=None):
        if time.time() - self._token_at > 1800:
            self.refresh()
        h = {"Authorization": "Bearer " + self._token}
        h.update(extra or {})
        return h

    def request(self, method, url, ok=(200,), retries=4, **kw):
        """One Graph call with retries on 429 / 5xx and network errors. Returns the response (any status in `ok`)."""
        kw.setdefault("timeout", 120)
        extra = kw.pop("headers", None)
        last = None
        for attempt in range(retries):
            try:
                r = requests.request(method, url, headers=self.headers(extra), **kw)
            except requests.exceptions.RequestException as e:
                last = e
                time.sleep(3 * (attempt + 1))
                continue
            if r.status_code == 401 and attempt == 0:
                self.refresh()
                continue
            if r.status_code in (429, 502, 503, 504):
                try:
                    wait = int(r.headers.get("Retry-After", 0))
                except ValueError:
                    wait = 0
                time.sleep(max(wait, 5 * (attempt + 1)))
                last = r
                continue
            return r
        if isinstance(last, requests.Response):
            return last
        raise RuntimeError(f"network error: {type(last).__name__ if last else 'unknown'}")

    def live_item(self, share_url=None):
        share_url = share_url or ADMIN_SHARE_URL
        variants = [share_url]
        if "&wdwpf=" in share_url:                       # the same link without the tracking parameter
            variants.append(share_url.split("&wdwpf=")[0])
        code = None
        for u in variants:
            b64 = base64.b64encode(u.encode()).decode().rstrip("=").replace("/", "_").replace("+", "-")
            r = self.request("GET", f"{GRAPH}/shares/u!{b64}/driveItem")
            if r.status_code == 200:
                return r.json()
            code = r.status_code
        die(f"cannot find the SharePoint file (HTTP {code}). Check the share link in the script.")

    def download(self, item, dest):
        url = item.get("@microsoft.graph.downloadUrl")
        if url:
            r = requests.get(url, stream=True, timeout=300)
        else:
            r = self.request("GET", f"{GRAPH}/drives/{item['parentReference']['driveId']}/items/{item['id']}/content",
                             stream=True)
        if r.status_code != 200:
            die(f"cannot download the live file (HTTP {r.status_code}).")
        with open(dest, "wb") as f:
            for b in r.iter_content(1 << 20):
                f.write(b)
        return dest


def same_content(a, b):
    """True when two xlsx files hold the same workbook. SharePoint rewrites the document properties of an Office file it
    stores, so the bytes can differ by a few KB although the data are the same; then the sheet data parts are compared."""
    if sha256_file(a) == sha256_file(b):
        return True, "identical bytes"
    try:
        def parts(path):
            with zipfile.ZipFile(path) as z:
                return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()
                        if n.startswith("xl/worksheets/") or n in ("xl/sharedStrings.xml", "xl/styles.xml")}
        pa, pb = parts(a), parts(b)
        if pa and pa == pb:
            return True, "same sheet data (SharePoint rewrote the document properties)"
    except (zipfile.BadZipFile, OSError):
        pass
    return False, ""


# ----------------------------------------------------------------------------------------------- backup tool
def run_backup(repo, key=BACKUP_KEY):
    script = Path(repo) / "sp_versions.py"
    if not script.is_file():
        return False, f"{script} not found"
    try:
        p = subprocess.run([sys.executable, str(script), "backup", key], cwd=str(repo), capture_output=True,
                           text=True, timeout=900)
    except Exception as e:
        return False, f"{type(e).__name__}: {e}"
    tail = (p.stdout or "").strip().splitlines()[-2:]
    return p.returncode == 0, " | ".join(tail) if tail else (p.stderr or "").strip()[-200:]


def backup_holds(repo, sha, key=BACKUP_KEY):
    cat = Path(repo) / "catalog" / "versions.csv"
    try:
        with open(cat, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                if row.get("key") == key and (row.get("sha256") or "").lower() == sha.lower():
                    snap = Path(repo) / (row.get("snapshot") or "")
                    return row.get("version"), snap.is_file()
    except OSError:
        pass
    return None, False


# ----------------------------------------------------------------------------------------------- upload
def upload_replace(g, item, local, etag, name=TARGET_NAME, flag="uploaded"):
    """Replace the content of the existing item through an upload session. If-Match makes SharePoint refuse the
    upload when the file changed since we looked at it (HTTP 412)."""
    d, i = item["parentReference"]["driveId"], item["id"]
    r = g.request("POST", f"{GRAPH}/drives/{d}/items/{i}/createUploadSession",
                  headers={"If-Match": etag} if etag else None,
                  json={"item": {"@microsoft.graph.conflictBehavior": "replace", "name": name}})
    if r.status_code == 412:
        die("somebody changed the file on SharePoint after the checks. Nothing was replaced. Run the script again.")
    if r.status_code in (423, 409):
        die(f"SharePoint says the file is locked (someone has {name} open for editing). "
            "Ask them to close it, then run the script again.")
    if r.status_code not in (200, 201):
        die(f"SharePoint refused the upload (HTTP {r.status_code}): {r.text[:200]}")
    url = r.json()["uploadUrl"]
    size = Path(local).stat().st_size
    start, tries = 0, 0
    with open(local, "rb") as f:
        while start < size:
            f.seek(start)
            chunk = f.read(CHUNK)
            end = start + len(chunk) - 1
            try:
                rr = requests.put(url, data=chunk, timeout=300,
                                  headers={"Content-Length": str(len(chunk)),
                                           "Content-Range": f"bytes {start}-{end}/{size}"})
                code = rr.status_code
            except requests.exceptions.RequestException:
                rr, code = None, 0
            if code in (200, 201, 202):
                start += len(chunk)
                tries = 0
                log(f"   uploaded {human(min(start, size))} of {human(size)}")
                continue
            if tries < 4 and code in (0, 416, 429, 500, 502, 503, 504):
                tries += 1
                time.sleep(5 * tries)
                try:                                  # where does the session expect the next byte?
                    st = requests.get(url, timeout=60).json()
                    rng = st.get("nextExpectedRanges") or []
                    if rng:
                        start = int(str(rng[0]).split("-")[0])
                except Exception:
                    pass
                continue
            try:
                requests.delete(url, timeout=30)      # drop the half-open session
            except Exception:
                pass
            die(f"the upload failed (HTTP {code}). The file on SharePoint was not replaced.")
    STATE[flag] = True


def item_now(g, item):
    r = g.request("GET", f"{GRAPH}/drives/{item['parentReference']['driveId']}/items/{item['id']}")
    return r.json() if r.status_code == 200 else {}


# ----------------------------------------------------------------------------------------------- storage
def delete_versions(g, item):
    """Delete every historical version of the file. Returns (deleted, freed_bytes, failed)."""
    d, i = item["parentReference"]["driveId"], item["id"]
    url = f"{GRAPH}/drives/{d}/items/{i}/versions"
    versions = []
    while url:
        r = g.request("GET", url)
        if r.status_code != 200:
            return 0, 0, -1
        body = r.json()
        versions += body.get("value", [])
        url = body.get("@odata.nextLink")
    versions.sort(key=lambda v: v.get("lastModifiedDateTime", ""), reverse=True)
    deleted = freed = failed = 0
    for n, v in enumerate(versions):
        r = g.request("DELETE", f"{GRAPH}/drives/{d}/items/{i}/versions/{v['id']}")
        if r.status_code in (200, 204, 404):
            deleted += 1
            freed += int(v.get("size") or 0)
        elif n == 0 and r.status_code in (400, 403, 405):
            continue                                   # the newest entry is the current version: it cannot go
        else:
            failed += 1
    return deleted, freed, failed


def purge_recycle_bin(g, item, since_iso, name=TARGET_NAME):
    """Best effort: remove from the site recycle bin what the version clean-up just put there (same file name, deleted
    since the start of this run). Uses the Graph beta recycle bin API; needs a permission the app may not have."""
    try:
        site_id = item["parentReference"].get("siteId")
        if not site_id:
            return None, "the file has no site id"
        ids, freed = [], 0
        url = f"{GRAPH_BETA}/sites/{site_id}/recycleBin/items?$top=200"
        pages = 0
        while url and pages < 30:
            r = g.request("GET", url)
            if r.status_code != 200:
                return None, f"recycle bin not readable (HTTP {r.status_code})"
            pages += 1
            body = r.json()
            for it in body.get("value", []):
                if (it.get("name") or "").lower() == name.lower() and (it.get("deletedDateTime") or "") >= since_iso:
                    ids.append(it["id"])
                    freed += int(it.get("size") or 0)
            url = body.get("@odata.nextLink")
        if not ids:
            return 0, "nothing of this run found in the recycle bin"
        removed = 0
        for k in range(0, len(ids), 50):
            r = g.request("POST", f"{GRAPH_BETA}/sites/{site_id}/recycleBin/items/delete",
                          json={"ids": ids[k:k + 50]})
            if r.status_code not in (200, 204):
                return removed, f"recycle bin delete refused (HTTP {r.status_code}); {removed} removed"
            removed += len(ids[k:k + 50])
        return removed, f"{removed} item(s), about {human(freed)}"
    except Exception as e:
        return None, f"{type(e).__name__}"


# ----------------------------------------------------------------------------------------------- data dictionary
DICT_VERSIONED = re.compile(r"^AFRO_admin_data_dictionary_v(\d+)\.xlsx$", re.I)


def find_dictionary(src=None, dirs=None):
    """The dictionary that goes with the data file: next to it (AFRO_admin_data_dictionary_vNN.xlsx for the same NN,
    or the plain name, also in out/upload_to_sharepoint), else the newest one in the usual output folders."""
    folders = []
    if src is not None:
        src = Path(src)
        m = re.search(r"_v(\d+)", src.stem)
        if m:
            folders.append((src.parent, f"AFRO_admin_data_dictionary_v{m.group(1)}.xlsx"))
        folders += [(src.parent, DICT_NAME), (src.parent / "out" / "upload_to_sharepoint", DICT_NAME)]
        for d, n in folders:
            if (d / n).is_file():
                return d / n
    found = []
    for d in (dirs or DEFAULT_DIRS):
        d = Path(d)
        if d.is_dir():
            for p in d.iterdir():
                m = DICT_VERSIONED.match(p.name)
                if m:
                    found.append((int(m.group(1)), p))
                elif p.name.lower() == DICT_NAME.lower():
                    found.append((-1, p))
    return max(found, key=lambda t: (t[0], t[1].stat().st_mtime))[1] if found else None


def inspect_dictionary(path):
    """Sheets and the documented column names of a dictionary workbook."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        info = {"sheets": list(wb.sheetnames), "columns": [], "entries": 0}
        if DICT_SHEET in wb.sheetnames:
            rows = wb[DICT_SHEET].iter_rows(values_only=True)
            head = [("" if v is None else str(v)).strip() for v in next(rows, [])]
            ci = head.index("Column") if "Column" in head else None
            info["head"] = head
            if ci is not None:
                for r in rows:
                    if len(r) > ci and r[ci] not in (None, ""):
                        info["columns"].append(str(r[ci]).strip())
            info["entries"] = len(info["columns"])
        return info
    finally:
        wb.close()


def content_key(path):
    """A hash of the cell content of a workbook that does not change when SharePoint rewrites document properties."""
    h = hashlib.sha256()
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if n.startswith("xl/worksheets/") or n in ("xl/sharedStrings.xml", "xl/workbook.xml"):
                h.update(n.encode())
                h.update(z.read(n))
    return h.hexdigest()


def read_json_file(p, default):
    try:
        return json.loads(Path(p).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


def publish_dictionary(g, args, tmp, started_iso, dict_path, data_header=None, dry=False):
    """Replace the SharePoint data dictionary with dict_path. Returns (status, message); status is one of
    'published', 'unchanged', 'planned' (dry run), 'failed'. Never exits: the admin data may already be published."""
    STATE["soft"] = True
    try:
        log("")
        log("-" * 78)
        log("Data dictionary" + ("  [DRY RUN]" if dry else ""))
        dict_path = Path(dict_path)
        if not dict_path.is_file():
            die(f"the dictionary file {dict_path} does not exist.")
        log(f"Dictionary file : {dict_path}  ({human(dict_path.stat().st_size)})")
        try:
            new = inspect_dictionary(dict_path)
        except Exception as e:
            die(f"the dictionary cannot be read ({type(e).__name__}). Is it a valid .xlsx?")
        problems = []
        if DICT_SHEET not in new["sheets"]:
            problems.append(f"no sheet '{DICT_SHEET}' (sheets: {', '.join(new['sheets'])})")
        elif new["entries"] < 50:
            problems.append(f"only {new['entries']} columns documented - the dictionary looks incomplete")
        if problems:
            die("the dictionary did not pass the checks:\n   - " + "\n   - ".join(problems))
        undocumented = [c for c in (data_header or []) if c and c not in new["columns"]]
        missing_req = [c for c in REQUIRED_COLUMNS if c not in new["columns"]]
        log(f"                  {len(new['sheets'])} sheets, {new['entries']} documented columns")
        if undocumented:
            log(f"   WARNING: {len(undocumented)} column(s) of the data are not in the dictionary: "
                + ", ".join(undocumented[:12]) + (" ..." if len(undocumented) > 12 else ""))
        elif missing_req:
            log("   WARNING: key columns not documented: " + ", ".join(missing_req))

        item = g.live_item(DICT_SHARE_URL)
        if (item.get("name") or "").lower() != DICT_NAME.lower():
            die(f"the file found through the dictionary share link is called '{item.get('name')}', not {DICT_NAME}.")
        etag = item.get("eTag") or item.get("cTag")
        mod_by = ((item.get("lastModifiedBy") or {}).get("user") or {}).get("displayName", "unknown")
        log(f"On SharePoint   : {item['name']}, {human(item.get('size', 0))}, last modified "
            f"{item.get('lastModifiedDateTime', '?')} by {mod_by}")
        live_path = g.download(item, tmp / "dict_live.xlsx")
        live_sha = sha256_file(live_path)
        try:
            live = inspect_dictionary(live_path)
        except Exception:
            live = None
        if live:
            log(f"Live dictionary : {len(live['sheets'])} sheets, {live['entries']} documented columns")
            lost_sheets = [x for x in live["sheets"] if x not in new["sheets"]]
            if lost_sheets and not args.force:
                die("the new dictionary lacks sheets that the live one has: " + ", ".join(lost_sheets)
                    + ". Use --force if that is intended.")
            if live["entries"] and new["entries"] < 0.7 * live["entries"] and not args.force:
                die(f"the new dictionary documents {new['entries']} columns, the live one {live['entries']}. "
                    f"Use --force if that is intended.")
            gone = [c for c in live["columns"] if c not in new["columns"]]
            came = [c for c in new["columns"] if c not in live["columns"]]
            if came:
                log(f"   new in the dictionary    ({len(came)}): " + ", ".join(came[:10]) + (" ..." if len(came) > 10 else ""))
            if gone:
                log(f"   no longer in it          ({len(gone)}): " + ", ".join(gone[:10]) + (" ..." if len(gone) > 10 else ""))
        same, why = same_content(live_path, dict_path)
        bdir = Path(args.dictionary_backup_dir).expanduser()
        state_file = bdir / "dictionary_state.json"
        st = read_json_file(state_file, {})
        edited = False
        if st.get("content_key"):
            try:
                edited = content_key(live_path) != st["content_key"]
            except Exception:
                edited = False
        if edited:
            log(f"   NOTE: the live dictionary differs from the one this tool pushed last time ({st.get('at', '?')}): "
                f"someone changed it on SharePoint. Its copy is kept in {bdir}.")
        if same and not args.force:
            log(f"The SharePoint dictionary already holds this version ({why}); it will not be uploaded again.")
            return "unchanged", "already up to date"
        log("Plan: " + ("(dry run) " if dry else "") + f"keep a local copy of the live dictionary in {bdir}; "
            f"replace {DICT_NAME} on SharePoint; check it; "
            + ("keep" if args.keep_versions else "delete") + " the SharePoint version history.")
        if dry:
            return "planned", "dictionary passes every check"

        # local copy of what is live now (the backup tool does not track the dictionary)
        bdir.mkdir(parents=True, exist_ok=True)
        stamp = (item.get("lastModifiedDateTime") or "").replace("-", "").replace(":", "").replace("T", "_").rstrip("Z")
        keep = bdir / f"AFRO_admin_data_dictionary_{stamp or 'live'}_{live_sha[:8]}.xlsx"
        if not any(live_sha[:8] in f.name for f in bdir.glob("AFRO_admin_data_dictionary_*.xlsx")):
            shutil.copyfile(live_path, keep)
            if sha256_file(keep) != live_sha:
                die("the local copy of the live dictionary could not be verified, so the upload was not started.")
            log(f"Local copy of the live dictionary: {keep.name}")
        else:
            log("The live dictionary is already in the local copies.")

        staged = tmp / DICT_NAME
        shutil.copyfile(dict_path, staged)
        log(f"Uploading {DICT_NAME} ({human(staged.stat().st_size)}) ...")
        upload_replace(g, item, staged, etag, name=DICT_NAME, flag="dict_uploaded")
        now = item_now(g, item)
        got, want = int(now.get("size", -1)), staged.stat().st_size
        if abs(got - want) > max(0.02 * want, 16384):
            die(f"after the upload SharePoint reports {got} bytes, expected about {want}. The previous dictionary is in {bdir}.")
        back = g.download(now if now.get("@microsoft.graph.downloadUrl") else item, tmp / "dict_check.xlsx")
        ok_same, why = same_content(back, staged)
        if not ok_same:
            try:
                x, y = inspect_dictionary(back), inspect_dictionary(staged)
                if x["sheets"] == y["sheets"] and x["columns"] == y["columns"]:
                    ok_same, why = True, "same sheets and documented columns (file repackaged by SharePoint)"
            except Exception:
                pass
        if not ok_same:
            die(f"the dictionary on SharePoint is not the uploaded one. The previous dictionary is in {bdir}.")
        log(f"Upload checked: SharePoint holds the new dictionary ({why}).")
        try:
            write = {"content_key": content_key(back), "sha256": sha256_file(back), "at": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                     "source": dict_path.name}
            state_file.write_text(json.dumps(write, indent=1), encoding="utf-8")
        except Exception:
            pass
        if args.keep_versions:
            log("Dictionary version history kept (--keep-versions).")
        else:
            deleted, freed, failed = delete_versions(g, item)
            if failed == -1:
                log("WARNING: could not read the dictionary version history.")
            else:
                log(f"Dictionary version history: {deleted} old version(s) deleted"
                    + (f", {failed} could not be deleted" if failed else ""))
                if deleted and not args.no_purge_bin:
                    n, info = purge_recycle_bin(g, item, started_iso, name=DICT_NAME)
                    log("Recycle bin: " + (info if n is not None else f"not emptied ({info})"))
        log(f"DONE. {DICT_NAME} on SharePoint is {dict_path.name}.")
        return "published", dict_path.name
    except SoftStop as e:
        log("")
        log("DICTIONARY NOT PUSHED - " + str(e))
        if STATE.get("dict_uploaded"):
            log("The upload had already happened, see above.")
        else:
            log("The dictionary on SharePoint was not changed.")
        return "failed", str(e).splitlines()[0]
    finally:
        STATE["soft"] = False


# ----------------------------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Publish the latest corrected admin data to SharePoint "
                                             "(replaces AFRO_admin_data.xlsx, keeps no versions there).")
    ap.add_argument("--file", help="the corrected file to publish (default: the highest AFRO_admin_data_vNN.xlsx)")
    ap.add_argument("--dir", help="folder to look in for AFRO_admin_data_vNN.xlsx")
    ap.add_argument("--dry-run", action="store_true", help="run every check, write nothing")
    ap.add_argument("--yes", action="store_true", help="do not ask for confirmation")
    ap.add_argument("--secrets", help="a secrets.env with the three SHAREPOINT_* values")
    ap.add_argument("--backup-repo", default=str(DEFAULT_BACKUP_REPO), help="folder of the SharePoint_version_backup tool")
    ap.add_argument("--no-backup", action="store_true",
                    help="skip the local backup of the live file (NOT recommended; the old file would be lost)")
    ap.add_argument("--keep-versions", action="store_true", help="do not delete the SharePoint version history")
    ap.add_argument("--no-purge-bin", action="store_true", help="do not try to empty the deleted copies from the recycle bin")
    ap.add_argument("--allow-unsaved", action="store_true", help="accept a file whose check_dp values are empty")
    ap.add_argument("--max-row-drop", type=float, default=MAX_ROW_DROP,
                    help="largest accepted drop in rows versus the live file (default 0.10 = 10%%)")
    ap.add_argument("--force", action="store_true", help="accept a row drop or an identical file (use with care)")
    ap.add_argument("--dictionary", help="the data dictionary to push (default: the one next to the data file)")
    ap.add_argument("--no-dictionary", action="store_true", help="do not push the data dictionary")
    ap.add_argument("--dictionary-only", action="store_true", help="push only the data dictionary, not the admin data")
    ap.add_argument("--dictionary-backup-dir", default=str(HERE / "dictionary_backups"),
                    help="where the previous live dictionary is kept before it is replaced")
    ap.add_argument("--result-json", help="write what happened (admin / dictionary status) to this file")
    args = ap.parse_args()
    args.backup_repo = str(Path(args.backup_repo).expanduser().resolve())

    started = dt.datetime.now(dt.timezone.utc)
    started_iso = started.strftime("%Y-%m-%dT%H:%M:%SZ")
    result = {"admin": "not run", "dictionary": "not run", "dictionary_message": "", "started": started_iso}

    def save_result():
        if args.result_json:
            try:
                Path(args.result_json).write_text(json.dumps(result, indent=1), encoding="utf-8")
            except OSError:
                pass
    log("=" * 78)
    log("Publish the admin data to SharePoint" + ("  [DRY RUN - nothing will be written]" if args.dry_run else ""))

    if args.dictionary_only:
        dp = Path(args.dictionary) if args.dictionary else find_dictionary()
        if not dp or not dp.is_file():
            die("no dictionary file found: give it with --dictionary FILE.")
        g = Graph(args.secrets)
        tmp = Path(tempfile.mkdtemp(prefix="publish_dict_"))
        try:
            status, msg = publish_dictionary(g, args, tmp, started_iso, dp, dry=args.dry_run)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        result.update({"admin": "skipped", "dictionary": status, "dictionary_message": msg})
        save_result()
        return EXIT_DICT_FAILED if status == "failed" else 0

    # 1. the file
    src, cands = choose_file(args)
    log(f"File to publish : {src}")
    if len(cands) > 1:
        log("Other candidates: " + ", ".join(c[2].name for c in cands if c[2] != src))
    log(f"                  {human(src.stat().st_size)}, saved {dt.datetime.fromtimestamp(src.stat().st_mtime):%Y-%m-%d %H:%M}")
    log("Reading the file ...")
    new = inspect_workbook(src)
    problems, warnings = check_new_file(src, new)
    if args.allow_unsaved:
        problems = [p for p in problems if "check_dp" not in p]
    log(f"                  {new['rows']:,} records, {len(new['header'])} columns, {new.get('countries', '?')} countries")
    for w in warnings:
        log("   WARNING: " + w)
    if problems:
        die("the file did not pass the checks:\n   - " + "\n   - ".join(problems))
    dict_path = None
    if not args.no_dictionary:
        dict_path = Path(args.dictionary) if args.dictionary else find_dictionary(src)
        if args.dictionary and not dict_path.is_file():
            die(f"the dictionary file {dict_path} does not exist.")
        if dict_path:
            log(f"Dictionary      : {dict_path}")
        else:
            log("Dictionary      : none found next to the data file - it will not be pushed (give it with --dictionary FILE)")

    # 2. the live file
    g = Graph(args.secrets)
    item = g.live_item()
    etag = item.get("eTag") or item.get("cTag")
    live_name = item.get("name", "")
    if live_name.lower() != TARGET_NAME.lower():
        die(f"the file found through the share link is called '{live_name}', not {TARGET_NAME}.")
    mod_by = ((item.get("lastModifiedBy") or {}).get("user") or {}).get("displayName", "unknown")
    log(f"On SharePoint   : {live_name}, {human(item.get('size', 0))}, last modified "
        f"{item.get('lastModifiedDateTime', '?')} by {mod_by}")
    tmp = Path(tempfile.mkdtemp(prefix="publish_admin_"))
    try:
        live_path = g.download(item, tmp / "live.xlsx")
        live_sha = sha256_file(live_path)
        already, why = same_content(live_path, src)
        if already and not args.force:
            log(f"The file on SharePoint already holds this version ({why}); it will not be uploaded again.")
        else:
            live = inspect_workbook(live_path)
            d_rows = new["rows"] - live["rows"]
            pct = d_rows / live["rows"] if live["rows"] else 0
            log(f"Live file now   : {live['rows']:,} records, {len(live['header'])} columns")
            log(f"Change          : {d_rows:+,} records ({pct:+.1%})")
            lost = [h for h in live["header"] if h and h not in new["header"]]
            gained = [h for h in new["header"] if h and h not in live["header"]]
            if gained:
                log("   new columns: " + ", ".join(gained))
            if lost:
                die("the new file lacks columns that the live file has: " + ", ".join(lost[:15]))
            if -pct > args.max_row_drop and not args.force:
                die(f"the new file has {-pct:.1%} fewer records than the live file (limit {args.max_row_drop:.0%}). "
                    f"Check the file, or use --force if this is expected.")
        skip_upload = already and not args.force
        # somebody may have changed the live file after this file was made: their edits would be overwritten
        try:
            live_mod = dt.datetime.fromisoformat(item.get("lastModifiedDateTime", "").replace("Z", "+00:00"))
            src_mod = dt.datetime.fromtimestamp(src.stat().st_mtime, dt.timezone.utc)
        except ValueError:
            live_mod = src_mod = None
        if not skip_upload and live_mod and live_mod > src_mod:
            log("")
            log(f"WARNING: the file on SharePoint was modified by {mod_by} at {item.get('lastModifiedDateTime')},")
            log(f"         AFTER {src.name} was saved ({src_mod:%Y-%m-%d %H:%M} UTC). Publishing would overwrite that change.")
            log("         It is saved in the local backup, but check what changed first (python sp_versions.py compare admin N M)")
            log("         and rebuild from the newest version if needed.")
            if not args.force:
                die("the live file is newer than the file to publish. Re-run with --force once you checked it.")

        log("")
        log("Plan:")
        if skip_upload:
            log("   1. (no upload: SharePoint already holds this version)")
            log("   2. (no upload)")
        else:
            log(f"   1. back up the live file locally ({args.backup_repo})" if not args.no_backup else "   1. (local backup skipped)")
            log(f"   2. replace AFRO_admin_data.xlsx on SharePoint with {src.name} (uploaded as {TARGET_NAME})")
        log("   3. delete the SharePoint version history of the file" + (" (skipped)" if args.keep_versions else ""))
        log("   4. back up again, so the new file is the newest local version")
        if args.dry_run:
            if dict_path:
                publish_dictionary(g, args, tmp, started_iso, dict_path, data_header=new["header"], dry=True)
            log("")
            log("Dry run finished: the file passes every check. Run again without --dry-run to publish.")
            return 0
        if not args.yes:
            ans = input("\nType YES to publish: ").strip()
            if ans != "YES":
                log("Cancelled. Nothing was changed.")
                return 0

        if not skip_upload:
            # 3. local backup of the live file
            if not args.no_backup:
                log("Backing up the live file ...")
                ok, msg = run_backup(args.backup_repo)
                log("   " + msg)
                ver, present = backup_holds(args.backup_repo, live_sha)
                if not ok or ver is None or not present:
                    die("the local backup does not hold the file that is on SharePoint now, so the upload was not started. "
                        "Run run_backup.bat in SharePoint_version_backup, or use --no-backup if you accept the risk.")
                log(f"   the live file is saved locally as backup version {ver}")

            # 4. upload
            staged = tmp / TARGET_NAME
            shutil.copyfile(src, staged)              # the copy carries the SharePoint name; the original keeps its name
            log(f"Uploading {TARGET_NAME} ({human(staged.stat().st_size)}) ...")
            upload_replace(g, item, staged, etag)

            # 5. verify
            now = item_now(g, item)
            got = int(now.get("size", -1))
            if abs(got - staged.stat().st_size) > 0.02 * staged.stat().st_size:
                die(f"after the upload SharePoint reports {got} bytes, expected about {staged.stat().st_size}. "
                    f"Restore the previous file from the local backup (restore {BACKUP_KEY} latest).")
            back = g.download(now if now.get("@microsoft.graph.downloadUrl") else item, tmp / "check.xlsx")
            ok_same, why = same_content(back, staged)
            if not ok_same:
                x, y = inspect_workbook(back), inspect_workbook(staged)
                if x["header"] == y["header"] and x["rows"] == y["rows"]:
                    ok_same, why = True, "same columns and record count (file repackaged by SharePoint)"
                    log("   WARNING: the sheet data are not byte-identical to the uploaded file; same columns and records.")
            if not ok_same:
                die("the file on SharePoint is not the uploaded one. Restore the previous file from the local backup.")
            log(f"Upload checked: SharePoint holds the new file ({why}).")

        result["admin"] = "unchanged" if skip_upload else "published"
        save_result()

        # 6. storage
        if args.keep_versions:
            log("Version history kept (--keep-versions).")
        else:
            deleted, freed, failed = delete_versions(g, item)
            if failed == -1:
                log("WARNING: could not read the version history; run again later or clean it in SharePoint.")
            else:
                log(f"Version history: {deleted} old version(s) deleted"
                    + (f" (about {human(freed)})" if freed else "") + (f", {failed} could not be deleted" if failed else ""))
                if deleted and not args.no_purge_bin:
                    n, info = purge_recycle_bin(g, item, started_iso)
                    if n is None:
                        log(f"Recycle bin: not emptied ({info}). The deleted copies still use space until the site recycle "
                            f"bin is emptied (Site contents > Recycle bin > Second-stage) or ask IT.")
                    else:
                        log(f"Recycle bin: {info}")

        # 7. backup of the new live file
        if not args.no_backup:
            ok, msg = run_backup(args.backup_repo)
            log("Local backup after the upload: " + (msg or ("done" if ok else "failed")))

        log("")
        log("DONE. AFRO_admin_data.xlsx on SharePoint is " + src.name + ".")

        # 8. the data dictionary
        rc = 0
        if dict_path:
            status, msg = publish_dictionary(g, args, tmp, started_iso, dict_path, data_header=new["header"])
            result.update({"dictionary": status, "dictionary_message": msg})
            if status == "failed":
                rc = EXIT_DICT_FAILED
        else:
            result["dictionary"] = "skipped" if args.no_dictionary else "none found"
        save_result()
        log("")
        if rc:
            log("The admin data is published, but the dictionary is not: fix the cause above, then run "
                "'python admin.py dictionary' (or publish_admin_data.py --dictionary-only --dictionary FILE).")
        else:
            log("Next: tell the countries/focal points if needed, and send the feedback workbook.")
        return rc
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        log("Cancelled.")
        sys.exit(130)
